import io
import re
import zipfile
from pathlib import PurePath
from pypdf import PdfReader
from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation

TYPES = {
    "pdf": {"application/pdf"},
    "docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    "pptx": {
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    },
}


def extract(data: bytes, filename: str, mime: str):
    filename = re.sub(r"[^\w. ()-]", "_", PurePath(filename.replace("\\", "/")).name)[
        :200
    ]
    extension = filename.rsplit(".", 1)[-1].lower()
    if extension not in TYPES or mime not in TYPES.get(extension, set()) | {
        "application/octet-stream",
        "",
    }:
        raise ValueError(
            "This file type is not supported. Please upload PDF, DOCX, or PPTX."
        )
    if extension == "pdf":
        if not data.startswith(b"%PDF-"):
            raise ValueError("The file is not a valid PDF.")
    else:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 5000
                or sum(e.file_size for e in entries) > 60 * 1024 * 1024
            ):
                raise ValueError(
                    "The document expands beyond the safe processing limit."
                )
            expected = (
                "word/document.xml" if extension == "docx" else "ppt/presentation.xml"
            )
            if expected not in archive.namelist():
                raise ValueError("File contents do not match the extension.")
    sections = []
    if extension == "pdf":
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs are not supported.")
        if len(reader.pages) > 500:
            raise ValueError("Please split PDFs longer than 500 pages.")
        for i, page in enumerate(reader.pages, 1):
            text = (
                (page.extract_text(extraction_mode="layout") or "")
                if "/Contents" in page
                else ""
            )
            sections.append(
                {
                    "page": i,
                    "title": next(
                        (l.strip()[:150] for l in text.splitlines() if l.strip()),
                        f"Page {i}",
                    ),
                    "content": text,
                }
            )
    elif extension == "docx":
        document = Document(io.BytesIO(data))
        title, lines = "Document", []
        for block in document.iter_inner_content():
            if hasattr(block, "rows"):
                lines.extend(
                    " | ".join(c.text for c in row.cells) for row in block.rows
                )
            else:
                style = block.style.name if block.style else ""
                if style.startswith("Heading"):
                    if lines:
                        sections.append(
                            {
                                "page": len(sections) + 1,
                                "title": title,
                                "content": "\n".join(lines),
                            }
                        )
                    title, lines = block.text, []
                prefix = (
                    "- "
                    if "List" in style
                    or block._p.find(".//" + qn("w:numPr")) is not None
                    else ""
                )
                lines.append(prefix + block.text)
        if lines:
            sections.append(
                {"page": len(sections) + 1, "title": title, "content": "\n".join(lines)}
            )
    else:
        presentation = Presentation(io.BytesIO(data))
        if len(presentation.slides) > 500:
            raise ValueError("Please split presentations longer than 500 slides.")

        def shape_text(shapes):
            lines = []
            for shape in shapes:
                if hasattr(shape, "shapes"):
                    lines.extend(shape_text(shape.shapes))
                if shape.has_text_frame:
                    lines.extend(
                        ("  " * p.level) + ("- " if p.level else "") + p.text
                        for p in shape.text_frame.paragraphs
                    )
                if shape.has_table:
                    lines.extend(
                        " | ".join(c.text for c in row.cells)
                        for row in shape.table.rows
                    )
            return lines

        for i, slide in enumerate(presentation.slides, 1):
            lines = shape_text(slide.shapes)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                lines.append(
                    "Speaker notes: " + slide.notes_slide.notes_text_frame.text
                )
            sections.append(
                {
                    "page": i,
                    "title": (
                        slide.shapes.title.text if slide.shapes.title else f"Slide {i}"
                    ),
                    "content": "\n".join(lines),
                }
            )
    sections = [s for s in sections if s["content"].strip()]
    if not sections:
        raise ValueError(
            "No readable text was found in this document. Scanned images require OCR before upload."
        )
    if sum(len(s["content"]) for s in sections) > 400000:
        raise ValueError(
            "Extracted text exceeds 400,000 characters. Please split this document."
        )
    return {"source_type": extension, "filename": filename, "sections": sections}

import io
import json
from pathlib import Path
from unittest.mock import patch, AsyncMock
import pytest
from fastapi.testclient import TestClient
from docx import Document
from pptx import Presentation
from pypdf import PdfWriter
from main import app, calls
from services.extraction import extract
from services.gemini import parse_json, AIError
from services.generator import chunks, validate_sources
from utils.validation import Material, GenerateRequest, Quiz

client = TestClient(app)
@pytest.fixture(autouse=True)
def reset_limiter():
    calls.clear()
@pytest.fixture
def demo():
    return json.loads((Path(__file__).resolve().parents[2] / 'frontend/assets/demo.json').read_text())
def docx_data():
    doc = Document();doc.add_heading('Cell Biology',1);doc.add_paragraph('Cells are the basic units of life.');doc.add_paragraph('Nucleus stores DNA.','List Bullet')
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Ribosome';table.cell(0,1).text='Protein synthesis'
    stream=io.BytesIO();doc.save(stream);return stream.getvalue()
def request_data():
    return {'files':[{'filename':'notes.docx','source_type':'docx','sections':[{'page':1,'title':'Cells','content':'Cells are the basic units of life.'}]}],'material_types':['reviewer']}
def test_health():
    assert client.get('/health').json()['data']['status']=='ok'
    assert client.get('/').json()['success']
def test_docx_extract():
    result=extract(docx_data(),'../../notes.docx','application/octet-stream')
    assert result['filename']=='notes.docx'
    assert result['sections'][0]['title']=='Cell Biology'
    assert 'Ribosome | Protein synthesis' in result['sections'][0]['content']
    assert '- Nucleus' in result['sections'][0]['content']
def test_pptx_extract():
    pres=Presentation();slide=pres.slides.add_slide(pres.slide_layouts[1]);slide.shapes.title.text='Cells';slide.placeholders[1].text='Basic unit of life';slide.notes_slide.notes_text_frame.text='Remember cell theory'
    stream=io.BytesIO();pres.save(stream)
    result=extract(stream.getvalue(),'notes.pptx','application/octet-stream')
    assert result['sections'][0]['page']==1
    assert 'Remember cell theory' in result['sections'][0]['content']
def test_blank_pdf():
    writer=PdfWriter();writer.add_blank_page(width=100,height=100);stream=io.BytesIO();writer.write(stream)
    with pytest.raises(ValueError,match='No readable text'):extract(stream.getvalue(),'blank.pdf','application/pdf')
def test_bad_type_and_signature():
    with pytest.raises(ValueError):extract(b'hello','bad.exe','application/octet-stream')
    with pytest.raises(ValueError):extract(b'hello','bad.pdf','application/pdf')
    with pytest.raises(ValueError):extract(docx_data(),'bad.docx','application/pdf')
def test_upload_api():
    response=client.post('/api/extract',files=[('files',('notes.docx',docx_data(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'))])
    assert response.status_code==200
    assert response.json()['data']['characters']>0
def test_corrupt_controlled():
    result=client.post('/api/extract',files=[('files',('bad.docx',b'corrupt','application/octet-stream'))])
    assert result.status_code==400
    assert result.json()['error']['code']=='INVALID_FILE'
def test_file_limit():
    result=client.post('/api/extract',files=[('files',('large.pdf',b'%PDF-'+b'x'*(10*1024*1024),'application/pdf'))])
    assert result.status_code==413
def test_json_recovery():
    assert parse_json('```json\n{"a":1}\n```')=={'a':1}
    assert parse_json('Accidental prefix {"a":{"b":"}"}} suffix')=={'a':{'b':'}'}}
    with pytest.raises(ValueError):parse_json('no object')
    with pytest.raises(ValueError):parse_json('{"a":eval(1)}')
def test_schema(demo):
    assert Material.model_validate(demo).flashcards
    bad=demo['quiz'][0].copy();bad['correct_answer']='Not a choice'
    with pytest.raises(ValueError):Quiz.model_validate(bad)
def test_input_validation():
    payload=request_data();payload['question_count']=11
    result=client.post('/api/generate',json=payload)
    assert result.status_code==422
    assert not result.json()['success']
def test_large_chunking():
    payload=request_data();payload['files'][0]['sections'][0]['content']='A long paragraph. '*10000
    output=chunks(GenerateRequest.model_validate(payload).files)
    assert len(output)>1
    assert all(len(part)<=18000 for part in output)
    assert all('SOURCE notes.docx' in part for part in output)
def test_invalid_reference(demo):
    with pytest.raises(AIError):validate_sources(Material.model_validate(demo),GenerateRequest.model_validate(request_data()).files)
def test_generate_mock(demo):
    with patch('main.generate',new=AsyncMock(return_value=Material.model_validate(demo))):
        result=client.post('/api/generate',json=request_data())
    assert result.status_code==200
    assert len(result.json()['data']['flashcards'])==10
def test_unconfigured_ai(monkeypatch):
    monkeypatch.delenv('GEMINI_API_KEY',raising=False)
    result=client.post('/api/generate',json=request_data())
    assert result.status_code==502
    assert result.json()['error']['code']=='AI_ERROR'
def test_rate_limit(monkeypatch):
    monkeypatch.setenv('RATE_LIMIT_PER_HOUR','1')
    client.post('/api/generate',json={})
    assert client.post('/api/generate',json={}).status_code==429
def test_cors():
    response=client.options('/api/generate',headers={'Origin':'http://localhost:8080','Access-Control-Request-Method':'POST','Access-Control-Request-Headers':'content-type'})
    assert response.headers['access-control-allow-origin']=='http://localhost:8080'
def test_docs():
    assert client.get('/docs').status_code==200
    assert client.get('/openapi.json').status_code==200

def test_pdf_text():
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer=PdfWriter();page=writer.add_blank_page(width=300,height=300)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 30 260 Td (Cells are the basic units of life.) Tj ET')
    page[NameObject('/Contents')]=writer._add_object(stream)
    data=io.BytesIO();writer.write(data)
    result=extract(data.getvalue(),'cells.pdf','application/pdf')
    assert 'basic units' in result['sections'][0]['content']

def test_duplicate_source():
    payload=request_data();payload['files']*=2
    assert client.post('/api/generate',json=payload).status_code==422

def test_strict_boolean(demo):
    demo['true_false'][0]['answer']='yes'
    with pytest.raises(ValueError):Material.model_validate(demo)

def test_rate_error_has_cors(monkeypatch):
    monkeypatch.setenv('RATE_LIMIT_PER_HOUR','0')
    result=client.post('/api/generate',json={},headers={'Origin':'http://localhost:8080'})
    assert result.status_code==429
    assert result.headers['access-control-allow-origin']=='http://localhost:8080'

def test_gemini_repair(monkeypatch, demo):
    import asyncio
    from services.gemini import generate_json
    monkeypatch.setenv('GEMINI_API_KEY','test-not-a-real-key');monkeypatch.setenv('GEMINI_MODEL','test-model')
    class Response:
        status_code=200
        is_error=False
        def __init__(self,text):self.text=text
        def json(self):return {'candidates':[{'content':{'parts':[{'text':self.text}]}}]}
    responses=[Response('{bad'),Response(json.dumps(demo))]
    mock=AsyncMock();mock.__aenter__.return_value=mock;mock.post.side_effect=responses
    with patch('services.gemini.httpx.AsyncClient',return_value=mock):
        result=asyncio.run(generate_json('source',Material))
    assert result.title==demo['title']
    assert mock.post.await_count==2

def test_gemini_invalid_controlled(monkeypatch):
    import asyncio
    from services.gemini import generate_json
    monkeypatch.setenv('GEMINI_API_KEY','test-not-a-real-key');monkeypatch.setenv('GEMINI_MODEL','test-model')
    mock=AsyncMock();mock.__aenter__.return_value=mock
    class Response:
        status_code=200
        is_error=False
        def json(self):return {'candidates':[{'content':{'parts':[{'text':'not json'}]}}]}
    mock.post.return_value=Response()
    with patch('services.gemini.httpx.AsyncClient',return_value=mock):
        with pytest.raises(AIError,match='could not be validated'):asyncio.run(generate_json('source',Material))
    assert mock.post.await_count==3
